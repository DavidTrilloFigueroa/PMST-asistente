"""Datos sintéticos de prueba, sin contenido de la documentación privada."""

from copy import deepcopy
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from src.importar_catalogo import (
    ASOCIACIONES, CAMPOS_ASOCIACION, CAMPOS_MEDIDA, CAMPOS_RELACION,
    CAMPOS_TIPO, MEDIDAS, RELACIONES, construir_catalogos, guardar_archivos,
    huellas, importar, validar_objetivos,
)


def informe_vacio():
    return {"errores": [], "advertencias": [], "datos_ausentes": []}


def datos():
    fuente = {"archivo": "sintetico.xlsx", "hoja": "Prueba", "fila": 2}
    medidas = []
    for letra in ("A", "B"):
        medida = dict.fromkeys(CAMPOS_MEDIDA)
        medida.update(id=f"MED-TEST-{letra}", nombre=f"Nombre de prueba {letra}",
                      objetivo="  Objetivo de prueba\nSegunda línea  ",
                      descripcion="Descripción sintética\n\nSin datos privados.",
                      fuente=deepcopy(fuente))
        medidas.append(medida)
    tipos = [{"tipo": "TIPO_SIMETRICO", "significado": "Definición de prueba",
              "direccion": "Simétrica", "fuente": deepcopy(fuente)},
             {"tipo": "TIPO_DIRIGIDO", "significado": "Definición de prueba",
              "direccion": "A → B", "fuente": deepcopy(fuente)}]
    relaciones = [{"origen": medidas[0]["id"], "nombre_origen": medidas[0]["nombre"],
                   "destino": medidas[1]["id"], "nombre_destino": medidas[1]["nombre"],
                   "tipo": "TIPO_SIMETRICO", "fuente": deepcopy(fuente)}]
    asociaciones = [{"id": "REL__MED-TEST-A__KPI-TEST-X", "medida": medidas[0]["id"],
                     "nombre_medida": medidas[0]["nombre"], "kpi": "KPI-TEST-X",
                     "nombre_kpi": "Indicador de prueba", "tipo": "TIPO_PRUEBA",
                     "subtipo": "SUBTIPO_PRUEBA", "fuente": deepcopy(fuente)}]
    return medidas, relaciones, tipos, asociaciones


def crear_excel(carpeta):
    medidas, relaciones, tipos, asociaciones = datos()
    tablas = [(MEDIDAS, [("Hoja1", CAMPOS_MEDIDA, medidas)]),
              (RELACIONES, [("Relaciones candidatas", CAMPOS_RELACION, relaciones),
                            ("Leyenda relaciones", CAMPOS_TIPO, tipos)]),
              (ASOCIACIONES, [("FICHA_MEDIDA_KPI", CAMPOS_ASOCIACION, asociaciones)])]
    for archivo, hojas in tablas:
        libro = Workbook()
        libro.remove(libro.active)
        for nombre, campos, registros in hojas:
            hoja = libro.create_sheet(nombre)
            # Orden distinto del esquema: obliga a leer por nombre de columna.
            claves = list(reversed(campos))
            hoja.append([campos[k] for k in claves])
            for registro in registros:
                hoja.append([registro[k] for k in claves])
            hoja.cell(20, 15).number_format = "0"  # Celdas vacías con formato.
        libro.save(carpeta / archivo)
        libro.close()


class ValidacionTests(unittest.TestCase):
    def setUp(self):
        self.medidas, self.relaciones, self.tipos, self.asociaciones = datos()
        self.informe = informe_vacio()

    def construir(self):
        return construir_catalogos(self.medidas, self.relaciones, self.tipos,
                                   self.asociaciones, self.informe)

    def codigos(self):
        return {e["codigo"] for e in self.informe["errores"]}

    def test_separacion_estricta_objetivos_y_texto(self):
        catalogos = self.construir()
        self.assertFalse(self.informe["errores"])
        self.assertEqual(catalogos["objetivos_busqueda.json"], [
            {"id": m["id"], "objetivo": m["objetivo"]} for m in self.medidas])
        self.assertFalse(validar_objetivos([{"id": "X", "objetivo": "Y", "nombre": "Z"}]))
        self.assertFalse(validar_objetivos([{"id": "X"}]))

    def test_referencias_invalidas_en_ambos_excel(self):
        self.relaciones[0]["destino"] = "NO-EXISTE"
        self.asociaciones[0]["medida"] = "NO-EXISTE"
        self.construir()
        self.assertEqual(sum(e["codigo"] == "referencia_invalida" for e in self.informe["errores"]), 2)

    def test_ids_duplicados_medidas_y_asociaciones(self):
        self.medidas.append(deepcopy(self.medidas[0]))
        self.asociaciones.append(deepcopy(self.asociaciones[0]))
        self.construir()
        self.assertEqual(sum(e["codigo"] == "id_duplicado" for e in self.informe["errores"]), 2)
        self.assertIn("asociacion_duplicada", self.codigos())

    def test_relacion_simetrica_inversa_duplicada(self):
        inversa = deepcopy(self.relaciones[0])
        inversa["origen"], inversa["destino"] = inversa["destino"], inversa["origen"]
        inversa["nombre_origen"], inversa["nombre_destino"] = inversa["nombre_destino"], inversa["nombre_origen"]
        self.relaciones.append(inversa)
        self.construir()
        self.assertIn("relacion_duplicada", self.codigos())

    def test_relaciones_dirigidas_no_se_confunden_ni_generan(self):
        self.relaciones[0]["tipo"] = "TIPO_DIRIGIDO"
        inversa = deepcopy(self.relaciones[0])
        inversa["origen"], inversa["destino"] = inversa["destino"], inversa["origen"]
        inversa["nombre_origen"], inversa["nombre_destino"] = inversa["nombre_destino"], inversa["nombre_origen"]
        self.relaciones.append(inversa)
        catalogos = self.construir()
        self.assertFalse(self.informe["errores"])
        self.assertEqual(catalogos["relaciones_medidas.json"], self.relaciones)

    def test_nombres_comparados_sin_cambiar_original(self):
        self.asociaciones[0]["nombre_medida"] = "  NOMBRE  DE\n PRUEBA A "
        self.construir()
        self.assertFalse(self.informe["errores"])
        self.assertEqual(self.asociaciones[0]["nombre_medida"], "  NOMBRE  DE\n PRUEBA A ")
        self.asociaciones[0]["nombre_medida"] = "Otro nombre"
        self.construir()
        self.assertIn("nombre_contradictorio", self.codigos())

    def test_id_asociacion_incoherente(self):
        self.asociaciones[0]["id"] = "REL__MED-TEST-B__KPI-TEST-X"
        self.construir()
        self.assertIn("id_asociacion_incoherente", self.codigos())

    def test_kpi_fuentes_y_tipos_por_asociacion(self):
        otra = deepcopy(self.asociaciones[0])
        otra.update(id="REL__MED-TEST-B__KPI-TEST-X", medida=self.medidas[1]["id"],
                    nombre_medida=self.medidas[1]["nombre"], tipo="OTRO_TIPO",
                    nombre_kpi="  INDICADOR DE PRUEBA  ")
        otra["fuente"]["fila"] = 3
        self.asociaciones.append(otra)
        catalogos = self.construir()
        self.assertFalse(self.informe["errores"])
        kpi = catalogos["kpis.json"][0]
        self.assertEqual(len(kpi["fuentes"]), 2)
        self.assertEqual(kpi["fuentes"][1]["nombre_original"], otra["nombre_kpi"])
        self.assertNotIn("tipo", kpi)
        self.assertNotIn("subtipo", kpi)
        otra["nombre_kpi"] = "Indicador contradictorio"
        self.construir()
        self.assertIn("nombre_contradictorio", self.codigos())

    def test_tipo_y_direccion_desconocidos(self):
        self.relaciones[0]["tipo"] = "NO_DEFINIDO"
        self.tipos[0]["direccion"] = "Sin definir"
        self.construir()
        self.assertTrue({"tipo_desconocido", "direccion_desconocida"} <= self.codigos())

    def test_obligatorios_y_ausencias_opcionales(self):
        self.construir()
        self.assertFalse(self.informe["errores"])
        self.assertEqual(len(self.informe["advertencias"]), 5)
        self.medidas[0]["objetivo"] = None
        self.medidas[1]["descripcion"] = "   "
        self.construir()
        self.assertEqual(sum(e["codigo"] == "campo_obligatorio" for e in self.informe["errores"]), 2)


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.temporal = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporal.cleanup)
        self.entrada = Path(self.temporal.name) / "entrada"
        self.salida = Path(self.temporal.name) / "salida"
        self.entrada.mkdir()
        crear_excel(self.entrada)

    def test_importacion_fiel_reordenacion_columnas_y_huellas(self):
        antes = huellas(self.entrada)
        informe = importar(self.entrada, self.salida)
        self.assertEqual(informe["estado"], "correcto")
        self.assertEqual(len(list(self.salida.glob("*.json"))), 7)
        medidas = json.loads((self.salida / "medidas.json").read_text(encoding="utf-8"))
        self.assertEqual(medidas[0]["objetivo"], datos()[0][0]["objetivo"])
        self.assertIsNone(medidas[0]["fases_implementacion"])
        self.assertEqual(medidas[0]["fuente"]["fila"], 2)
        self.assertEqual(medidas[0]["fuente"]["columnas"]["id"], "I")
        self.assertEqual(antes, huellas(self.entrada))

    def test_error_no_sustituye_catalogo_previo(self):
        importar(self.entrada, self.salida)
        anteriores = {p.name: p.read_bytes() for p in self.salida.glob("*.json") if p.name != "informe_validacion.json"}
        archivo = self.entrada / MEDIDAS
        libro = load_workbook(archivo)
        libro["Hoja1"].cell(3, 9).value = "MED-TEST-A"
        libro.save(archivo)
        libro.close()
        informe = importar(self.entrada, self.salida)
        self.assertEqual(informe["estado"], "error")
        self.assertEqual(anteriores, {nombre: (self.salida / nombre).read_bytes() for nombre in anteriores})

    def test_fila_sin_id_y_contenido_sin_encabezado(self):
        archivo = self.entrada / MEDIDAS
        libro = load_workbook(archivo)
        libro["Hoja1"].cell(4, 10).value = "Contenido sin encabezado"
        libro.save(archivo)
        libro.close()
        informe = importar(self.entrada, self.salida)
        errores = [e for e in informe["errores"] if e["codigo"] == "id_ausente"]
        self.assertEqual(errores[0]["fuente"]["archivo"], MEDIDAS)
        self.assertEqual(errores[0]["fuente"]["hoja"], "Hoja1")
        self.assertEqual(errores[0]["fuente"]["fila"], 4)
        self.assertFalse((self.salida / "medidas.json").exists())

    def test_formula_no_se_importa_como_texto_ni_cache(self):
        archivo = self.entrada / MEDIDAS
        libro = load_workbook(archivo)
        libro["Hoja1"].cell(2, 7).value = '=CONCAT("a", "b")'
        libro.save(archivo)
        libro.close()
        informe = importar(self.entrada, self.salida)
        self.assertIn("formula_o_error_excel", {e["codigo"] for e in informe["errores"]})

    def test_fases_y_kpis_ficha_se_conservan_sin_crear_asociaciones(self):
        archivo = self.entrada / MEDIDAS
        libro = load_workbook(archivo)
        fases = "  2. Texto inicial\n\n1. Texto posterior  "
        ficha = "Texto sintético de ficha\nSin asociaciones documentadas"
        libro["Hoja1"].cell(3, 5).value = fases
        libro["Hoja1"].cell(3, 4).value = ficha
        libro.save(archivo)
        libro.close()
        informe = importar(self.entrada, self.salida)
        medidas = json.loads((self.salida / "medidas.json").read_text(encoding="utf-8"))
        self.assertEqual(medidas[1]["fases_implementacion"], fases)
        self.assertEqual(medidas[1]["kpis_texto_ficha"], ficha)
        self.assertEqual(informe["recuentos"]["medidas_sin_kpis"], 1)
        self.assertEqual(informe["recuentos"]["asociaciones_medida_kpi"], 1)

    def test_encabezado_requerido_ausente(self):
        archivo = self.entrada / MEDIDAS
        libro = load_workbook(archivo)
        libro["Hoja1"].cell(1, 9).value = "ENCABEZADO_INCORRECTO"
        libro.save(archivo)
        libro.close()
        informe = importar(self.entrada, self.salida)
        self.assertIn("columnas_ausentes", {e["codigo"] for e in informe["errores"]})
        self.assertFalse((self.salida / "medidas.json").exists())

    def test_archivo_ausente_genera_informe(self):
        (self.entrada / MEDIDAS).unlink()
        informe = importar(self.entrada, self.salida)
        self.assertEqual(informe["estado"], "error")
        self.assertTrue((self.salida / "informe_validacion.json").exists())

    def test_reversion_si_falla_publicacion(self):
        self.salida.mkdir()
        (self.salida / "a.json").write_bytes(b"anterior")
        reemplazar = os.replace
        llamadas = 0

        def fallo_segundo(origen, destino):
            nonlocal llamadas
            llamadas += 1
            if llamadas == 2:
                raise OSError("Fallo simulado de escritura")
            return reemplazar(origen, destino)

        with patch("src.importar_catalogo.os.replace", side_effect=fallo_segundo):
            with self.assertRaises(OSError):
                guardar_archivos(self.salida, {"a.json": [1], "b.json": [2]})
        self.assertEqual((self.salida / "a.json").read_bytes(), b"anterior")
        self.assertFalse((self.salida / "b.json").exists())


if __name__ == "__main__":
    unittest.main()
