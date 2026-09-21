"""Importa los Excel locales sin modificarlos ni recurrir a servicios externos."""

import argparse
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import sys
import tempfile
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException


RAIZ = Path(__file__).resolve().parents[1]
MEDIDAS = "MEDIDAS.xlsx"
RELACIONES = "Relaciones medidas-medidas.xlsx"
ASOCIACIONES = "Relaciones medidas-KPIs.xlsx"
CAMPOS_MEDIDA = {
    "id": "ID", "nombre": "Nombre medida", "objetivo": "Objetivo",
    "descripcion": "Descripción", "fases_implementacion": "Fases de implementación",
    "kpis_texto_ficha": "KPI's Éxito/Eficacia", "resultados": "Resultados",
    "presupuesto": "Presupuesto", "actores_implicados": "Actores implicados",
}
CAMPOS_RELACION = {
    "origen": "ID_medida_origen", "nombre_origen": "Nombre_medida_origen",
    "tipo": "Tipo_relacion", "destino": "ID_medida_destino",
    "nombre_destino": "Nombre_medida_destino",
}
CAMPOS_TIPO = {"tipo": "Relación", "significado": "Significado", "direccion": "Dirección"}
CAMPOS_ASOCIACION = {
    "id": "ID_RELACION", "medida": "ID_MEDIDA", "nombre_medida": "NOMBRE_MEDIDA",
    "kpi": "ID_KPI", "nombre_kpi": "NOMBRE_KPI",
    "tipo": "TIPO_RELACION_PRINCIPAL", "subtipo": "SUBTIPO_INDICADOR_PRINCIPAL",
}


def ausente(valor):
    return valor is None or (isinstance(valor, str) and not valor.strip())


def nombre_comparable(valor):
    """Solo para comparar nombres; nunca sustituye el texto guardado."""
    return " ".join(valor.split()).casefold() if isinstance(valor, str) else valor


def incidencia(informe, nivel, codigo, mensaje, fuente=None, **detalles):
    informe[nivel].append({"codigo": codigo, "mensaje": mensaje,
                           "fuente": fuente, **detalles})


def leer_tabla(carpeta, archivo, hoja, campos, ids, informe):
    """Lee por encabezados exactos y conserva texto, filas y columnas originales."""
    fuente_base = {"archivo": archivo, "hoja": hoja, "fila": 1}
    registros = []
    try:
        libro = load_workbook(carpeta / archivo, read_only=True, data_only=False)
        try:
            if hoja not in libro.sheetnames:
                incidencia(informe, "errores", "hoja_ausente", "No existe la hoja requerida.", fuente_base)
                return []
            filas = libro[hoja].iter_rows()
            encabezados = next(filas, ())
            posiciones = {}
            for i, celda in enumerate(encabezados):
                nombre = celda.value
                if ausente(nombre):
                    continue
                if nombre in posiciones:
                    incidencia(informe, "errores", "encabezado_duplicado", "Encabezado repetido.", fuente_base, encabezado=nombre)
                posiciones[nombre] = i
            faltantes = set(campos.values()) - posiciones.keys()
            if faltantes:
                incidencia(informe, "errores", "columnas_ausentes", "Faltan columnas requeridas.", fuente_base, columnas=sorted(faltantes))
                return []
            extras = posiciones.keys() - set(campos.values())
            if extras:
                incidencia(informe, "errores", "columnas_no_previstas", "Revisar columnas adicionales antes de importarlas.", fuente_base, columnas=list(extras))
            for numero, fila in enumerate(filas, 2):
                if all(ausente(c.value) for c in fila):
                    continue
                fuente = {"archivo": archivo, "hoja": hoja, "fila": numero,
                          "columnas": {k: get_column_letter(posiciones[v] + 1) for k, v in campos.items()}}
                for i, celda in enumerate(fila):
                    if not ausente(celda.value) and i not in posiciones.values():
                        incidencia(informe, "errores", "contenido_sin_encabezado", "Hay contenido en una columna sin encabezado.", fuente, columna=get_column_letter(i + 1))
                    if celda.data_type in ("f", "e"):
                        incidencia(informe, "errores", "formula_o_error_excel", "Revisar la fórmula o el error Excel; no se calcula ni se sustituye por un valor almacenado.", fuente, columna=get_column_letter(i + 1))
                registro = {k: fila[posiciones[v]].value for k, v in campos.items()}
                registro = {k: None if v == "" else v for k, v in registro.items()}
                registro["fuente"] = fuente
                for clave in ids:
                    valor = registro[clave]
                    if ausente(valor):
                        incidencia(informe, "errores", "id_ausente", "Fila con contenido pero sin identificador.", fuente, campo=clave)
                    elif not isinstance(valor, str):
                        incidencia(informe, "errores", "id_no_textual", "El identificador debe ser texto y se conserva sin normalizar.", fuente, campo=clave)
                registros.append(registro)
        finally:
            libro.close()
    except (OSError, BadZipFile, InvalidFileException, ValueError, KeyError) as exc:
        incidencia(informe, "errores", "lectura_excel", str(exc), fuente_base)
    return registros


def indice_unico(registros, clave, entidad, informe):
    indice = {}
    for registro in registros:
        valor = registro[clave]
        if not isinstance(valor, str) or ausente(valor):
            continue
        if valor in indice:
            incidencia(informe, "errores", "id_duplicado", f"Identificador duplicado de {entidad}.", registro["fuente"], id=valor, primera_fuente=indice[valor]["fuente"])
        else:
            indice[valor] = registro
    return indice


def comprobar_nombre(original, referencia, informe, fuente, campo):
    if not isinstance(referencia, str) or ausente(referencia):
        incidencia(informe, "errores", "nombre_ausente", "Falta un nombre textual para validar la identidad.", fuente, campo=campo)
    elif nombre_comparable(original) != nombre_comparable(referencia):
        incidencia(informe, "errores", "nombre_contradictorio", "El nombre no corresponde al identificador.", fuente, campo=campo)


def validar_objetivos(objetivos):
    return all(set(r) == {"id", "objetivo"} for r in objetivos)


def construir_catalogos(medidas, relaciones, tipos, asociaciones, informe):
    por_medida = indice_unico(medidas, "id", "medida", informe)
    por_tipo = indice_unico(tipos, "tipo", "tipo de relación", informe)
    indice_unico(asociaciones, "id", "asociación", informe)
    for medida in medidas:
        for campo in ("nombre", "objetivo", "descripcion"):
            if not isinstance(medida[campo], str) or ausente(medida[campo]):
                incidencia(informe, "errores", "campo_obligatorio", "Falta un campo textual obligatorio.", medida["fuente"], campo=campo, id=medida["id"])
        for campo in CAMPOS_MEDIDA:
            if ausente(medida[campo]):
                informe["datos_ausentes"].append({"entidad": "medida", "id": medida["id"], "campo": campo, "fuente": medida["fuente"]})
        for campo in ("fases_implementacion", "actores_implicados"):
            if ausente(medida[campo]):
                incidencia(informe, "advertencias", "dato_ausente", "No consta en la documentación", medida["fuente"], id=medida["id"], campo=campo)
    for tipo in tipos:
        if tipo["direccion"] not in ("Simétrica", "A → B"):
            incidencia(informe, "errores", "direccion_desconocida", "La dirección de la leyenda requiere revisión; no se infiere.", tipo["fuente"])
        if ausente(tipo["significado"]):
            incidencia(informe, "errores", "significado_ausente", "Falta el significado del tipo de relación.", tipo["fuente"])

    def referencia(registro, clave, nombre):
        medida = por_medida.get(registro[clave])
        if medida is None:
            incidencia(informe, "errores", "referencia_invalida", "La medida referenciada no existe.", registro["fuente"], campo=clave, id=registro[clave])
        else:
            comprobar_nombre(medida["nombre"], registro[nombre], informe, registro["fuente"], nombre)

    vistas = set()
    for relacion in relaciones:
        referencia(relacion, "origen", "nombre_origen")
        referencia(relacion, "destino", "nombre_destino")
        tipo = por_tipo.get(relacion["tipo"])
        if tipo is None:
            incidencia(informe, "errores", "tipo_desconocido", "El tipo no figura en la leyenda.", relacion["fuente"], tipo=relacion["tipo"])
        extremos = (relacion["origen"], relacion["destino"])
        if tipo and tipo["direccion"] == "Simétrica":
            extremos = frozenset(extremos)
        clave = (relacion["tipo"], extremos)
        if clave in vistas:
            incidencia(informe, "errores", "relacion_duplicada", "Relación repetida, teniendo en cuenta la simetría de la leyenda.", relacion["fuente"])
        vistas.add(clave)

    kpis = {}
    pares = set()
    con_kpis = set()
    for asociacion in asociaciones:
        referencia(asociacion, "medida", "nombre_medida")
        esperado = f"REL__{asociacion['medida']}__{asociacion['kpi']}"
        if asociacion["id"] != esperado:
            incidencia(informe, "errores", "id_asociacion_incoherente", "ID_RELACION no coincide con REL__<ID_MEDIDA>__<ID_KPI>.", asociacion["fuente"], id=asociacion["id"])
        par = (asociacion["medida"], asociacion["kpi"])
        if par in pares:
            incidencia(informe, "errores", "asociacion_duplicada", "Pareja medida–KPI duplicada.", asociacion["fuente"])
        pares.add(par)
        con_kpis.add(asociacion["medida"])
        id_kpi = asociacion["kpi"]
        if not isinstance(id_kpi, str) or ausente(id_kpi):
            continue
        nombre = asociacion["nombre_kpi"]
        if not isinstance(nombre, str) or ausente(nombre):
            incidencia(informe, "errores", "nombre_kpi_ausente", "Falta el nombre del KPI.", asociacion["fuente"])
        if id_kpi not in kpis:
            kpis[id_kpi] = {"id": id_kpi, "nombre": nombre, "fuentes": []}
        else:
            comprobar_nombre(kpis[id_kpi]["nombre"], nombre, informe, asociacion["fuente"], "nombre_kpi")
        kpis[id_kpi]["fuentes"].append({**asociacion["fuente"], "nombre_original": nombre})
        for campo in ("tipo", "subtipo"):
            if ausente(asociacion[campo]):
                informe["datos_ausentes"].append({"entidad": "asociacion", "id": asociacion["id"], "campo": campo, "fuente": asociacion["fuente"]})
                incidencia(informe, "advertencias", "dato_ausente", "No consta en la documentación", asociacion["fuente"], id=asociacion["id"], campo=campo)
    sin_kpis = [m["id"] for m in medidas if m["id"] not in con_kpis]
    for id_medida in sin_kpis:
        fuente = por_medida[id_medida]["fuente"] if id_medida in por_medida else None
        incidencia(informe, "advertencias", "sin_kpis_asociados", "No consta en la documentación", fuente, id=id_medida, campo="kpis_asociados")
        informe["datos_ausentes"].append({"entidad": "medida", "id": id_medida, "campo": "kpis_asociados", "fuente": fuente})
    objetivos = [{"id": m["id"], "objetivo": m["objetivo"]} for m in medidas]
    if not validar_objetivos(objetivos):
        incidencia(informe, "errores", "catalogo_objetivos_invalido", "Solo se permiten id y objetivo.")
    informe["recuentos"] = {
        "medidas": len(medidas), "medidas_ids_unicos": len(por_medida),
        "relaciones_medidas": len(relaciones), "tipos_relacion": len(tipos),
        "asociaciones_medida_kpi": len(asociaciones), "kpis_distintos": len(kpis),
        "medidas_sin_fases": sum(ausente(m["fases_implementacion"]) for m in medidas),
        "medidas_sin_actores": sum(ausente(m["actores_implicados"]) for m in medidas),
        "medidas_sin_kpis": len(sin_kpis),
    }
    return {
        "medidas.json": medidas, "objetivos_busqueda.json": objetivos,
        "relaciones_medidas.json": relaciones, "tipos_relacion.json": tipos,
        "kpis.json": list(kpis.values()), "asociaciones_medida_kpi.json": asociaciones,
    }


def huellas(carpeta):
    return {nombre: hashlib.sha256((carpeta / nombre).read_bytes()).hexdigest()
            for nombre in (MEDIDAS, RELACIONES, ASOCIACIONES)}


def guardar_archivos(carpeta, documentos):
    """Serializa todo antes de reemplazar; revierte reemplazos ante errores de E/S.

    Cada reemplazo es atómico. No es una transacción frente a cortes eléctricos
    ni admite importaciones simultáneas: ejecutar una sola importación cada vez.
    """
    contenido = {nombre: (json.dumps(datos, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8")
                 for nombre, datos in documentos.items()}
    carpeta.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".importacion-", dir=carpeta) as temporal:
        etapa = Path(temporal)
        anteriores = {}
        for nombre, datos in contenido.items():
            destino = carpeta / nombre
            anteriores[nombre] = destino.read_bytes() if destino.exists() else None
            (etapa / nombre).write_bytes(datos)
        sustituidos = []
        try:
            for nombre in contenido:
                os.replace(etapa / nombre, carpeta / nombre)
                sustituidos.append(nombre)
        except OSError:
            for nombre in reversed(sustituidos):
                destino = carpeta / nombre
                if anteriores[nombre] is None:
                    destino.unlink()
                else:
                    respaldo = etapa / nombre
                    respaldo.write_bytes(anteriores[nombre])
                    os.replace(respaldo, destino)
            raise


def importar(entrada, salida):
    entrada, salida = Path(entrada).resolve(), Path(salida).resolve()
    if entrada == salida or entrada in salida.parents:
        raise ValueError("La salida no puede estar dentro de los documentos originales.")
    informe = {"fecha_utc": datetime.now(timezone.utc).isoformat(),
               "estado": "error", "recuentos": {}, "errores": [],
               "advertencias": [], "datos_ausentes": []}
    try:
        iniciales = huellas(entrada)
        medidas = leer_tabla(entrada, MEDIDAS, "Hoja1", CAMPOS_MEDIDA, ("id",), informe)
        relaciones = leer_tabla(entrada, RELACIONES, "Relaciones candidatas", CAMPOS_RELACION, ("origen", "destino"), informe)
        tipos = leer_tabla(entrada, RELACIONES, "Leyenda relaciones", CAMPOS_TIPO, ("tipo",), informe)
        asociaciones = leer_tabla(entrada, ASOCIACIONES, "FICHA_MEDIDA_KPI", CAMPOS_ASOCIACION, ("id", "medida", "kpi"), informe)
        catalogos = construir_catalogos(medidas, relaciones, tipos, asociaciones, informe)
        if not medidas:
            incidencia(informe, "errores", "catalogo_vacio", "No se ha leído ninguna medida.")
        finales = huellas(entrada)
        informe["integridad_originales"] = {"sha256_antes": iniciales, "sha256_despues": finales, "sin_cambios": iniciales == finales}
        if iniciales != finales:
            incidencia(informe, "errores", "originales_cambiados", "Los originales cambiaron durante la importación.")
        if not informe["errores"]:
            informe["estado"] = "correcto"
            guardar_archivos(salida, {**catalogos, "informe_validacion.json": informe})
            return informe
    except (OSError, BadZipFile, ValueError, TypeError) as exc:
        incidencia(informe, "errores", "importacion_interrumpida", str(exc))
    informe["estado"] = "error"
    guardar_archivos(salida, {"informe_validacion.json": informe})
    return informe


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--entrada", type=Path, default=RAIZ / "data/originales")
    parser.add_argument("--salida", type=Path, default=RAIZ / "data/procesados")
    args = parser.parse_args()
    try:
        informe = importar(args.entrada, args.salida)
    except (OSError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    if informe["errores"]:
        print("ERROR: importación rechazada. No se publica un nuevo catálogo. Consulte informe_validacion.json.", file=sys.stderr)
        for error in informe["errores"]:
            fuente = error.get("fuente") or {}
            print(f"- {error['codigo']}: {error['mensaje']} {fuente.get('archivo', '')} / {fuente.get('hoja', '')} / fila {fuente.get('fila', '')}", file=sys.stderr)
        return 1
    print("Catálogo generado. Excel originales sin cambios.")
    print(json.dumps(informe["recuentos"], ensure_ascii=False, indent=2))
    print(f"Advertencias: {len(informe['advertencias'])}. Detalle en informe_validacion.json.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
